#!/usr/bin/env python3
"""Build a sendable 16:9 PowerPoint from the Pravadh Labs investor deck."""
from __future__ import annotations

from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "investor" / "PravadhLabs-Investor-Pitch-Deck-2026.pptx"
PHOTO = ROOT / "assets" / "images" / "dhairya-shah.png"
ISV_BADGE = ROOT / "assets" / "images" / "partners" / "salesforce-isv-partner.png"
DESKTOP = Path("/Users/dhairyashah/Desktop/PRAVADH LABS/Pravadh Labs Investor ")
DESKTOP_OUT = DESKTOP / "PravadhLabs-Investor-Pitch-Deck-2026.pptx"

INK = RGBColor(0x0B, 0x14, 0x2A)
INK_SOFT = RGBColor(0x14, 0x1F, 0x3D)
PAPER = RGBColor(0xF7, 0xF7, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0x23, 0x2F, 0x52)
LINE_LIGHT = RGBColor(0xD4, 0xD4, 0xD0)
MUTED = RGBColor(0x86, 0x8A, 0x95)
MUTED_LIGHT = RGBColor(0xA9, 0xAF, 0xC0)
MUTED_INK = RGBColor(0x5A, 0x60, 0x70)
BLUE = RGBColor(0x1E, 0x5C, 0xEE)
AMBER = RGBColor(0xFD, 0x9E, 0x01)
TEAL = RGBColor(0x3E, 0x7C, 0x74)
AMBER_INK = RGBColor(0x1A, 0x12, 0x00)

FONT_DISPLAY = "Georgia"
FONT_BODY = "Calibri"
FONT_MONO = "Consolas"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
ML = Inches(0.58)
MR = Inches(12.75)
CONTENT_W = Inches(12.17)


def rgb_hex(c: RGBColor) -> str:
    return f"{c[0]:02X}{c[1]:02X}{c[2]:02X}"


def set_bg(slide, color: RGBColor) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def no_line(shape) -> None:
    shape.line.fill.background()


def add_rect(slide, x, y, w, h, fill: RGBColor, line: RGBColor | None = None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        no_line(shape)
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(1)
    return shape


def add_round(slide, x, y, w, h, fill: RGBColor, line: RGBColor | None = None, adj=0.08):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    try:
        shape.adjustments[0] = adj
    except Exception:
        pass
    if line is None:
        no_line(shape)
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(1)
    return shape


def tb(slide, x, y, w, h):
    return slide.shapes.add_textbox(x, y, w, h)


def wipe_p(p):
    p.text = ""
    return p


def clear_tf(tf):
    tf.word_wrap = True
    try:
        tf.auto_size = None
    except Exception:
        pass
    return wipe_p(tf.paragraphs[0])


def run(p, text, *, size=14, bold=False, color=PAPER, font=FONT_BODY, italic=False):
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    r.font.name = font
    return r


def para(tf, text, *, size=14, bold=False, color=PAPER, font=FONT_BODY,
         align=PP_ALIGN.LEFT, space_after=6, space_before=0, first=False):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space_after)
    p.space_before = Pt(space_before)
    p.level = 0
    if text:
        run(p, text, size=size, bold=bold, color=color, font=font)
    return p


def set_anchor(tf, anchor=MSO_ANCHOR.TOP):
    tf.word_wrap = True
    try:
        tf.auto_size = None
    except Exception:
        pass
    tf._txBody.bodyPr.set("anchor", {
        MSO_ANCHOR.TOP: "t",
        MSO_ANCHOR.MIDDLE: "ctr",
        MSO_ANCHOR.BOTTOM: "b",
    }.get(anchor, "t"))


def add_bar(slide):
    add_rect(slide, 0, 0, Inches(9.2), Inches(0.07), BLUE)
    add_rect(slide, Inches(9.2), 0, Inches(4.133), Inches(0.07), AMBER)


def add_footer(slide, left="Pravadh Labs", right="Confidential", dark=True):
    color = MUTED if dark else MUTED
    box = tb(slide, ML, Inches(7.12), Inches(7.4), Inches(0.28))
    tf = box.text_frame
    tf.word_wrap = False
    p = clear_tf(tf)
    run(p, left, size=11, color=color, font=FONT_BODY)
    box2 = tb(slide, Inches(8.4), Inches(7.12), Inches(4.35), Inches(0.28))
    tf2 = box2.text_frame
    tf2.word_wrap = False
    p2 = clear_tf(tf2)
    p2.alignment = PP_ALIGN.RIGHT
    run(p2, right, size=11, color=color, font=FONT_BODY)


def add_tag(slide, text, num, dark=True):
    tag_c = AMBER if dark else BLUE
    box = tb(slide, ML, Inches(0.28), Inches(9.2), Inches(0.32))
    tf = box.text_frame
    p = clear_tf(tf)
    run(p, text.upper(), size=11, color=tag_c, font=FONT_MONO)
    nbox = tb(slide, Inches(11.15), Inches(0.28), Inches(1.6), Inches(0.32))
    ntf = nbox.text_frame
    np = clear_tf(ntf)
    np.alignment = PP_ALIGN.RIGHT
    run(np, num, size=12, color=MUTED, font=FONT_MONO)


def add_title(slide, text, *, dark=True, size=30, y=Inches(0.62), w=None, h=Inches(0.95)):
    box = tb(slide, ML, y, w or CONTENT_W, h)
    tf = box.text_frame
    tf.word_wrap = True
    p = clear_tf(tf)
    p.space_after = Pt(0)
    run(p, text, size=size, color=PAPER if dark else INK, font=FONT_DISPLAY)
    return box


def add_lead(slide, text, *, dark=True, y=Inches(1.55), w=None, h=Inches(0.7), size=16):
    box = tb(slide, ML, y, w or Inches(11.4), h)
    tf = box.text_frame
    tf.word_wrap = True
    p = clear_tf(tf)
    p.space_after = Pt(0)
    run(p, text, size=size, color=MUTED_LIGHT if dark else MUTED_INK, font=FONT_BODY)
    return box


def add_logo(slide, dark=True):
    mark = add_round(slide, ML, Inches(0.28), Inches(0.42), Inches(0.42), BLUE, adj=0.18)
    tf = mark.text_frame
    set_anchor(tf, MSO_ANCHOR.MIDDLE)
    p = clear_tf(tf)
    p.alignment = PP_ALIGN.CENTER
    run(p, "P", size=18, bold=True, color=WHITE, font=FONT_DISPLAY)
    word = tb(slide, Inches(1.08), Inches(0.30), Inches(3.2), Inches(0.38))
    tf = word.text_frame
    p = clear_tf(tf)
    run(p, "Pravadh® Labs", size=18, color=PAPER if dark else INK, font=FONT_DISPLAY)


def add_card(slide, x, y, w, h, *, title, body, num=None, mods=None,
             accent=BLUE, dark=True, title_size=16, body_size=13):
    fill = INK_SOFT if dark else WHITE
    border = LINE if dark else LINE_LIGHT
    add_round(slide, x, y, w, h, fill, border, adj=0.06)
    add_rect(slide, x + Emu(40000), y, w - Emu(80000), Inches(0.045), accent)
    box = tb(slide, x + Inches(0.16), y + Inches(0.12), w - Inches(0.32), h - Inches(0.22))
    tf = box.text_frame
    tf.word_wrap = True
    set_anchor(tf, MSO_ANCHOR.TOP)
    first = True
    if num:
        p = para(tf, num.upper(), size=10, color=BLUE, font=FONT_MONO,
                 space_after=4, first=first)
        first = False
    p = para(tf, title, size=title_size, bold=True, color=PAPER if dark else INK,
             font=FONT_BODY, space_after=6, first=first)
    first = False
    para(tf, body, size=body_size, color=MUTED_LIGHT if dark else MUTED_INK,
         font=FONT_BODY, space_after=4)
    if mods:
        para(tf, mods.upper(), size=10, color=AMBER, font=FONT_MONO, space_before=6, space_after=0)
    return box


def add_stat(slide, x, y, w, h, num, label, dark=True):
    fill = INK_SOFT if dark else WHITE
    border = LINE if dark else LINE_LIGHT
    add_round(slide, x, y, w, h, fill, border, adj=0.08)
    box = tb(slide, x + Inches(0.16), y + Inches(0.12), w - Inches(0.28), h - Inches(0.2))
    tf = box.text_frame
    tf.word_wrap = True
    para(tf, num, size=22, color=AMBER if dark else BLUE, font=FONT_DISPLAY,
         space_after=4, first=True)
    para(tf, label, size=12, color=MUTED_LIGHT if dark else MUTED_INK, font=FONT_BODY, space_after=0)


def set_cell_borders(cell, color: RGBColor, width_pt=0.75):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag.endswith("}lnL") or child.tag.endswith("}lnR") or child.tag.endswith("}lnT") or child.tag.endswith("}lnB"):
            tcPr.remove(child)
    width = str(int(width_pt * 12700))
    hex_c = rgb_hex(color)
    for edge in ("lnL", "lnR", "lnT", "lnB"):
        ln = etree.SubElement(tcPr, qn(f"a:{edge}"))
        ln.set("w", width)
        ln.set("cap", "flat")
        ln.set("cmpd", "sng")
        ln.set("algn", "ctr")
        sf = etree.SubElement(ln, qn("a:solidFill"))
        srgb = etree.SubElement(sf, qn("a:srgbClr"))
        srgb.set("val", hex_c)
        etree.SubElement(ln, qn("a:prstDash")).set("val", "solid")


def fill_cell(cell, color: RGBColor):
    cell.fill.solid()
    cell.fill.fore_color.rgb = color


def write_cell(cell, text, *, size=12, bold=False, color=PAPER, font=FONT_BODY, align=PP_ALIGN.LEFT):
    tf = cell.text_frame
    tf.word_wrap = True
    set_anchor(tf, MSO_ANCHOR.MIDDLE)
    p = tf.paragraphs[0]
    p.alignment = align
    wipe_p(p)
    run(p, text, size=size, bold=bold, color=color, font=font)
    cell.margin_left = Inches(0.1)
    cell.margin_right = Inches(0.1)
    cell.margin_top = Inches(0.06)
    cell.margin_bottom = Inches(0.06)


def write_cell_runs(cell, parts, *, align=PP_ALIGN.LEFT):
    tf = cell.text_frame
    tf.word_wrap = True
    set_anchor(tf, MSO_ANCHOR.MIDDLE)
    p = tf.paragraphs[0]
    p.alignment = align
    wipe_p(p)
    for part in parts:
        run(p, part.get("text", ""), size=part.get("size", 12), bold=part.get("bold", False),
            color=part.get("color", PAPER), font=part.get("font", FONT_BODY))
    cell.margin_left = Inches(0.1)
    cell.margin_right = Inches(0.1)
    cell.margin_top = Inches(0.06)
    cell.margin_bottom = Inches(0.06)


def style_table(table, dark=True, header=True):
    border = LINE if dark else LINE_LIGHT
    body_fill = INK if dark else PAPER
    head_fill = INK_SOFT if dark else WHITE
    for i, row in enumerate(table.rows):
        for cell in row.cells:
            fill_cell(cell, head_fill if (header and i == 0) else body_fill)
            set_cell_borders(cell, border, 0.6)


def add_hyperlink(run_obj, url: str):
    run_obj.hyperlink.address = url


def new_slide(prs, dark=True):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, INK if dark else PAPER)
    add_bar(slide)
    return slide


def build():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    prs.core_properties.title = "Pravadh Labs — Investor Deck 2026"
    prs.core_properties.author = "Pravadh Labs Private Limited"
    prs.core_properties.subject = "Confidential"
    prs.core_properties.category = "Investor"

    # 01 COVER
    s = new_slide(prs, dark=True)
    add_logo(s, dark=True)
    tag = tb(s, Inches(8.35), Inches(0.34), Inches(4.4), Inches(0.3))
    p = clear_tf(tag.text_frame)
    p.alignment = PP_ALIGN.RIGHT
    run(p, "CONFIDENTIAL · INVESTOR DECK 2026", size=11, color=AMBER, font=FONT_MONO)

    eye = tb(s, ML, Inches(1.35), CONTENT_W, Inches(0.32))
    p = clear_tf(eye.text_frame)
    run(p, "THE PROBLEM IN ONE LINE", size=12, color=AMBER, font=FONT_MONO)

    h = tb(s, ML, Inches(1.7), Inches(12.1), Inches(1.55))
    tf = h.text_frame
    tf.word_wrap = True
    p = clear_tf(tf)
    p.space_after = Pt(4)
    run(p, "Enterprise work is split across systems.", size=32, color=PAPER, font=FONT_DISPLAY)
    p2 = tf.add_paragraph()
    p2.space_after = Pt(0)
    run(p2, "The next action lives in no one’s tool.", size=32, color=AMBER, font=FONT_DISPLAY)

    lead = tb(s, ML, Inches(3.4), Inches(11.6), Inches(1.05))
    tf = lead.text_frame
    tf.word_wrap = True
    p = clear_tf(tf)
    run(p, "Pravadh® is the control plane for teams who already run CRM, ERP, delivery, and sales. We unify ship, move, and govern on that stack — then queue one next action. Almost no manual work. They still approve before anything moves.",
        size=16, color=MUTED_LIGHT, font=FONT_BODY)

    stats = [
        ("Who", "People who own the CRM org — architects, RevOps, IT, SIs"),
        ("What", "Four live products on one rail — IQ, GovernX, PM, SI"),
        ("How", "Sit in the mess. Do not rip the stack out."),
        ("Ask", "₹3–5 Cr to scale GTM on product that already ships"),
    ]
    sw, gap = Inches(2.92), Inches(0.14)
    x = ML
    for num, label in stats:
        add_stat(s, x, Inches(4.6), sw, Inches(1.35), num, label, dark=True)
        x += sw + gap
    add_footer(s, "Pravadh Labs Private Limited", "pravadhlabs.com", dark=True)

    # 02 PROBLEM
    s = new_slide(prs, True)
    add_tag(s, "The problem we solve", "02 / 12", True)
    add_title(s, "A deploy can succeed and the release still fails.", size=28, h=Inches(0.7))
    add_lead(s, "The work is already happening. It is just not in one place — so nobody can see the next action, and nobody owns the trail.",
             y=Inches(1.38), h=Inches(0.7))
    cards = [
        (BLUE, "What breaks", "Three jobs. Three owners. No lock.",
         "Code ships in one tool. Records move in another. The audit sits in a spreadsheet. Delivery wears someone else’s name. The next move is in someone’s memory."),
        (AMBER, "What it costs", "Busywork, drift, and rework",
         "Skipped loads. Two owners on one issue. Policy found on Friday. Mapping knowledge walks out with people. Alerts arrive after the fact."),
        (BLUE, "What AI does not fix", "A chat beside the stack is another silo",
         "Agents are a given. If they cannot see deploy, data, and policy together, they cannot tell you what to do next — or wait for you to approve it."),
    ]
    cw, gap = Inches(3.94), Inches(0.16)
    x = ML
    for accent, num, title, body in cards:
        add_card(s, x, Inches(2.2), cw, Inches(4.5), num=num, title=title, body=body,
                 accent=accent, dark=True, title_size=17, body_size=14)
        x += cw + gap
    add_footer(s)

    # 03 CUSTOMERS
    s = new_slide(prs, False)
    add_tag(s, "Who buys", "03 / 12", False)
    add_title(s, "If you own the CRM org, you are the customer.", dark=False, size=28, h=Inches(0.62))
    add_lead(s, "One buyer set. Four jobs. The entry product is whichever job hurts this quarter — not a suite they have to swallow.",
             dark=False, y=Inches(1.28), h=Inches(0.5), size=15)
    buyers = [
        (BLUE, "01", "CRM architects", "Platform owners who still run deploy, migrate, and govern as three workstreams. They buy PravadhIQ."),
        (AMBER, "02", "RevOps & IT", "Teams who must keep CRM, ERP, and finance consistent — and explain drift. They buy GovernX."),
        (BLUE, "03", "Implementation SIs", "Consultancies delivering under their brand, on their domain or servers. They buy Pravadh PM and resell IQ."),
        (AMBER, "04", "Sales / RevOps leaders", "Need who, why, why now — written back to the CRM they already run. They buy Pravadh SI."),
    ]
    cw, ch, gap = Inches(5.98), Inches(1.85), Inches(0.18)
    positions = [(ML, Inches(1.9)), (ML + cw + gap, Inches(1.9)),
                 (ML, Inches(1.9 + ch + gap)), (ML + cw + gap, Inches(1.9 + ch + gap))]
    for (accent, num, title, body), (x, y) in zip(buyers, positions):
        add_card(s, x, y, cw, ch, num=num, title=title, body=body, accent=accent, dark=False)
    note = tb(s, ML, Inches(6.05), CONTENT_W, Inches(0.85))
    tf = note.text_frame
    tf.word_wrap = True
    p = clear_tf(tf)
    run(p, "Not a rip-and-replace buyer. Not a “new system of record” buyer. They keep CRM, ERP, finance, and delivery tools. They buy the process that sits between them.",
        size=14, color=MUTED_INK, font=FONT_BODY)
    add_footer(s, dark=False)

    # 04 HOW WE SOLVE
    s = new_slide(prs, True)
    add_tag(s, "How the product solves it", "04 / 12", True)
    add_title(s, "Each product is one hard job, completed. The rail is the expansion.",
              size=26, h=Inches(0.85))
    rows = [
        ["Customer", "The job that hurts", "What Pravadh does", "They still"],
        ["Architect / SI", "Ship, move, and govern as three tools", "IQ  One timeline — validate, map, audit, rollback", "Keep the CRM"],
        ["RevOps / data / finance", "CRM, ERP, and policy disagree", "GX  One pipeline + EMIE on the canvas", "Approve the fix"],
        ["Consultancy", "Delivery under someone else’s name", "PM  Their brand, domain, or private runtime", "Own the client"],
        ["Sales / RevOps", "Scores that never become an action", "SI  Who / why / why now — written back", "Run the CRM they have"],
    ]
    table_shape = s.shapes.add_table(5, 4, ML, Inches(1.62), CONTENT_W, Inches(4.35))
    table = table_shape.table
    table.columns[0].width = Inches(2.35)
    table.columns[1].width = Inches(3.55)
    table.columns[2].width = Inches(4.15)
    table.columns[3].width = Inches(2.12)
    style_table(table, dark=True)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = table.cell(r, c)
            if r == 0:
                write_cell(cell, val.upper(), size=10, bold=True, color=MUTED, font=FONT_MONO)
            elif c == 2:
                tag_txt, rest = val.split("  ", 1)
                write_cell_runs(cell, [
                    {"text": tag_txt + "  ", "size": 11, "bold": True, "color": AMBER, "font": FONT_MONO},
                    {"text": rest, "size": 13, "color": PAPER, "font": FONT_BODY},
                ])
            elif c == 0:
                write_cell(cell, val, size=13, color=MUTED_LIGHT, font=FONT_BODY)
            elif c == 3:
                write_cell(cell, val, size=13, bold=True, color=PAPER, font=FONT_BODY)
            else:
                write_cell(cell, val, size=13, color=MUTED_LIGHT, font=FONT_BODY)
    note = tb(s, ML, Inches(6.1), CONTENT_W, Inches(0.85))
    tf = note.text_frame
    tf.word_wrap = True
    p = clear_tf(tf)
    run(p, "Potential is not “more AI.” It is one process across jobs they already pay for — then a next action they can approve. That is what expands seats, modules, and products on the same account.",
        size=14, color=MUTED_LIGHT, font=FONT_BODY)
    add_footer(s)

    # 05 RAIL
    s = new_slide(prs, False)
    add_tag(s, "The rail — four products, one control plane", "05 / 12", False)
    add_title(s, "Depth-first products. Shared trail, approvals, and “you still approve.”",
              dark=False, size=26, h=Inches(0.85))
    add_round(s, ML, Inches(1.7), CONTENT_W, Inches(3.55), WHITE, LINE_LIGHT, adj=0.04)
    cap = tb(s, ML + Inches(0.28), Inches(1.86), Inches(8), Inches(0.28))
    p = clear_tf(cap.text_frame)
    run(p, "PRAVADH LABS CONTROL PLANE", size=11, color=MUTED, font=FONT_MONO)
    add_rect(s, ML + Inches(0.4), Inches(2.32), Inches(8.4), Inches(0.045), BLUE)
    add_rect(s, ML + Inches(8.8), Inches(2.32), Inches(2.9), Inches(0.045), AMBER)
    products = [
        ("PravadhIQ", "SHIP · MOVE · GOVERN", "CRM change", "https://pravadhiq.pravadhlabs.com"),
        ("GovernX", "EXTRACT · GOVERN · NOTIFY · EMIE", "Cross-platform data", "https://pravadhgx.pravadhlabs.com"),
        ("Pravadh PM", "ISSUES · VIEWS · BRAND · PRIVATE", "White-label delivery", "https://pravadhpm.pravadhlabs.com"),
        ("Pravadh SI", "CONNECT · QUALIFY · ACT", "Sales intelligence", "https://pravadhsi.pravadhlabs.com"),
    ]
    pw, gap = Inches(2.78), Inches(0.16)
    x = ML + Inches(0.28)
    for name, jobs, sub, url in products:
        add_round(s, x, Inches(2.62), pw, Inches(2.28), PAPER, LINE_LIGHT, adj=0.08)
        box = tb(s, x + Inches(0.1), Inches(2.74), pw - Inches(0.2), Inches(2.05))
        tf = box.text_frame
        tf.word_wrap = True
        p = para(tf, name, size=16, bold=True, color=INK, font=FONT_BODY, space_after=6,
                 align=PP_ALIGN.CENTER, first=True)
        add_hyperlink(p.runs[0], url)
        para(tf, jobs, size=11, color=MUTED_INK, font=FONT_MONO, space_after=6, align=PP_ALIGN.CENTER)
        para(tf, sub, size=12, color=MUTED, font=FONT_BODY, space_after=0, align=PP_ALIGN.CENTER)
        x += pw + gap

    pills = [
        (BLUE, "15+ sources into the CRM"),
        (AMBER, "Nothing moves without you"),
        (BLUE, "Stay on the stack you have"),
        (AMBER, "Audit trail by default"),
    ]
    x = ML
    for color, label in pills:
        w = Inches(2.9)
        shape = add_round(s, x, Inches(5.5), w, Inches(0.42), PAPER, color, adj=0.5)
        tf = shape.text_frame
        set_anchor(tf, MSO_ANCHOR.MIDDLE)
        p = clear_tf(tf)
        p.alignment = PP_ALIGN.CENTER
        run(p, label.upper(), size=10, color=color, font=FONT_MONO)
        x += w + Inches(0.14)
    add_footer(s, dark=False)

    # 06 WEDGE
    s = new_slide(prs, True)
    add_tag(s, "Why this can win — we don’t replace the stack", "06 / 12", True)
    add_title(s, "We sit in the mess between systems. That is the wedge.", size=26, h=Inches(0.7))
    add_lead(s, "Enterprises will not rip out CRM, ERP, finance, or delivery tooling for a new suite. They will buy a rail that makes the next action obvious on what they already run.",
             y=Inches(1.38), h=Inches(0.7), size=15)
    rows = [
        ["How the work usually sits", "How it sits on the rail"],
        ["Promoting code as its own job", "SHIP  Deploy with validation, dependencies, rollback"],
        ["Moving records as a second job", "MOVE  Map and transform before load"],
        ["The trail in a sheet", "GOVERN  One trail for every change and approval"],
        ["One-off extracts across CRM and ERP", "EXTRACT  One governed pipeline on the canvas"],
        ["The next move in someone’s memory", "EMIE / SI  Next action queued. You approve."],
        ["Delivery under someone else’s name", "PM  Your brand, your domain, or your servers"],
    ]
    table_shape = s.shapes.add_table(7, 2, ML, Inches(2.15), CONTENT_W, Inches(4.55))
    table = table_shape.table
    table.columns[0].width = Inches(5.4)
    table.columns[1].width = Inches(6.77)
    style_table(table, dark=True)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = table.cell(r, c)
            if r == 0:
                write_cell(cell, val.upper(), size=10, bold=True, color=MUTED, font=FONT_MONO)
            elif c == 0:
                write_cell(cell, val, size=13, color=MUTED_LIGHT, font=FONT_BODY)
            else:
                tag_txt, rest = val.split("  ", 1)
                write_cell_runs(cell, [
                    {"text": tag_txt + "  ", "size": 11, "bold": True, "color": AMBER, "font": FONT_MONO},
                    {"text": rest, "size": 13, "bold": True, "color": PAPER, "font": FONT_BODY},
                ])
    add_footer(s)

    # 07 LAND EXPAND
    s = new_slide(prs, False)
    add_tag(s, "Land on one job → expand on the rail", "07 / 12", False)
    left = tb(s, ML, Inches(0.7), Inches(6.35), Inches(5.8))
    tf = left.text_frame
    tf.word_wrap = True
    para(tf, "Same account. Next product when the next job shows up.",
         size=26, color=INK, font=FONT_DISPLAY, space_after=18, first=True)
    bullets = [
        ("Land IQ", " when ship / move / govern is still three jobs"),
        ("Attach GovernX", " when ERP and finance drift into the CRM"),
        ("Attach PM", " when the SI needs client-facing delivery under their name"),
        ("Attach SI", " when RevOps needs the next revenue action written back"),
    ]
    for strong, rest in bullets:
        p = tf.add_paragraph()
        p.space_after = Pt(12)
        run(p, "—  ", size=15, color=BLUE, font=FONT_BODY)
        run(p, strong, size=15, bold=True, color=INK, font=FONT_BODY)
        run(p, rest, size=15, color=MUTED_INK, font=FONT_BODY)
    add_card(s, Inches(7.2), Inches(1.55), Inches(5.55), Inches(2.2),
             title="Why expansion is natural",
             body="Same buyer, same org, same audit model. We are not cross-selling a second unrelated product. We are completing the process they already started.",
             accent=BLUE, dark=False, title_size=17, body_size=14)
    add_card(s, Inches(7.2), Inches(3.9), Inches(5.55), Inches(2.35),
             title="Why partners fit",
             body="SIs already sit in front of this buyer. They introduce. We deliver. They earn 20% of collected revenue while that customer pays — 5% one level down.",
             accent=AMBER, dark=False, title_size=17, body_size=14)
    add_footer(s, dark=False)

    # 08 PROOF
    s = new_slide(prs, True)
    add_tag(s, "Proof today", "08 / 12", True)
    left = tb(s, ML, Inches(0.7), Inches(6.45), Inches(5.9))
    tf = left.text_frame
    tf.word_wrap = True
    para(tf, "The studio is live. Commercial scale is the round.",
         size=26, color=PAPER, font=FONT_DISPLAY, space_after=16, first=True)
    proof = [
        "Four products in production — IQ, GovernX, PM, SI",
        "Salesforce ISV Partner",
        "15+ data sources connecting into the CRM",
        "Reseller program live — paid on collected cash, not bookings",
        "Dreamforce 2026 and West Africa Dreamin 2026",
        "SOC 2 Type II and ISO 27001 in progress",
    ]
    for item in proof:
        p = tf.add_paragraph()
        p.space_after = Pt(8)
        run(p, "—  ", size=15, color=AMBER, font=FONT_BODY)
        run(p, item, size=15, color=MUTED_LIGHT, font=FONT_BODY)
    p = tf.add_paragraph()
    p.space_before = Pt(10)
    run(p, "Revenue, pipeline, and ACV are shared in conversation and data room — not estimated on this slide.",
        size=13, color=MUTED, font=FONT_BODY)
    if ISV_BADGE.exists():
        s.shapes.add_picture(str(ISV_BADGE), Inches(7.2), Inches(0.78), Inches(1.18), Inches(1.18))
    add_card(s, Inches(7.2), Inches(2.12), Inches(5.55), Inches(2.05),
             num="What this proves", title="The problem is productised",
             body="Not a pitch for software we will build. Four complete jobs on one control plane, sold into the CRM org we already know how to reach.",
             accent=BLUE, dark=True, title_size=17, body_size=13)
    add_card(s, Inches(7.2), Inches(4.32), Inches(5.55), Inches(2.05),
             num="What the raise is for", title="Distribution, not invention",
             body="Partners, events, certifications, and the team to run more of the same motion — with a buyer we can name.",
             accent=AMBER, dark=True, title_size=17, body_size=13)
    add_footer(s)

    # 09 MODEL
    s = new_slide(prs, False)
    add_tag(s, "Business model", "09 / 12", False)
    add_title(s, "SaaS per product. Expand on the account. Partners on cash collected.",
              dark=False, size=24, h=Inches(0.75))
    models = [
        (BLUE, "PravadhIQ", "Land the CRM org. Services attach on complex migrations.", "Land · expand · migrate"),
        (AMBER, "GovernX", "Platform as systems connect. EMIE as the copilot on the canvas.", "Trial → platform → EMIE"),
        (BLUE, "Pravadh PM", "Seats. White-label and private runtime at a premium. Source stays with Labs.", "Cloud · brand · VPC"),
        (AMBER, "Pravadh SI", "Revenue motion on their CRM. Intro pricing through 31 Dec 2026.", "Connect · qualify · write back"),
    ]
    cw, ch, gap = Inches(5.98), Inches(1.72), Inches(0.18)
    positions = [(ML, Inches(1.5)), (ML + cw + gap, Inches(1.5)),
                 (ML, Inches(1.5 + ch + gap)), (ML + cw + gap, Inches(1.5 + ch + gap))]
    for (accent, title, body, mods), (x, y) in zip(models, positions):
        add_card(s, x, y, cw, ch, title=title, body=body, mods=mods, accent=accent, dark=False)
    stats = [
        ("20%", "Partner share of collected revenue, while the customer pays"),
        ("5%", "One-level referral — stops when that customer stops paying"),
        ("0", "Commission on unpaid invoices or verbal pipeline"),
        ("1 rail", "Cross-sell is the next job on the same org"),
    ]
    sw, gap = Inches(2.92), Inches(0.14)
    x = ML
    for num, label in stats:
        add_stat(s, x, Inches(5.3), sw, Inches(1.55), num, label, dark=False)
        x += sw + gap
    add_footer(s, dark=False)

    # 10 GTM
    s = new_slide(prs, True)
    add_tag(s, "Go-to-market", "10 / 12", True)
    add_title(s, "Reach the person who already owns the mess.", size=28, h=Inches(0.7))
    gtm = [
        (BLUE, "Where they already gather",
         "Dreamforce 2026. West Africa Dreamin 2026 sponsor. Live walkthroughs on real deploy, migrate, and governance scenarios — not booth scripts."),
        (AMBER, "Who already sits with them",
         "SIs introduce a qualified CRM buyer. We run product and delivery. They earn on cash we actually collect."),
        (BLUE, "How they try it",
         "GovernX and PM trials on the public products. SI intro pricing through year-end. IQ through real migration conversations."),
        (AMBER, "Where we build",
         "Founder-led across Ahmedabad, Bengaluru, and Jaipur. Cost-efficient product. Buyer access through the Salesforce ecosystem as an ISV partner."),
    ]
    cw, ch, gap = Inches(5.98), Inches(2.15), Inches(0.18)
    positions = [(ML, Inches(1.55)), (ML + cw + gap, Inches(1.55)),
                 (ML, Inches(1.55 + ch + gap)), (ML + cw + gap, Inches(1.55 + ch + gap))]
    for (accent, title, body), (x, y) in zip(gtm, positions):
        add_card(s, x, y, cw, ch, title=title, body=body, accent=accent, dark=True, title_size=17, body_size=14)
    add_footer(s)

    # 11 FOUNDER
    s = new_slide(prs, False)
    add_tag(s, "Founder / team advantage", "11 / 12", False)
    if PHOTO.exists():
        s.shapes.add_picture(str(PHOTO), ML, Inches(1.15), Inches(2.55), Inches(2.98))
        add_rect(s, ML, Inches(1.15), Inches(2.55), Inches(0.045), BLUE)
    name = tb(s, Inches(3.4), Inches(1.15), Inches(8.7), Inches(0.55))
    p = clear_tf(name.text_frame)
    run(p, "Dhairya Shah", size=32, color=INK, font=FONT_DISPLAY)
    role = tb(s, Inches(3.4), Inches(1.7), Inches(8.7), Inches(0.32))
    p = clear_tf(role.text_frame)
    run(p, "FOUNDER · ENTERPRISE / SALESFORCE ARCHITECT", size=12, color=MUTED_INK, font=FONT_MONO)
    creds = tb(s, Inches(3.4), Inches(2.15), Inches(8.7), Inches(2.2))
    tf = creds.text_frame
    tf.word_wrap = True
    items = [
        "11+ years in Salesforce and enterprise architecture",
        "Former VP — Salesforce Practice",
        "Agentblazer Legend",
        "Practitioner-led product development",
    ]
    first = True
    for item in items:
        p = para(tf, "", size=16, first=first, space_after=8)
        first = False
        run(p, "—  ", size=16, color=BLUE, font=FONT_BODY)
        run(p, item, size=16, color=MUTED_INK, font=FONT_BODY)
    bio = tb(s, ML, Inches(4.4), CONTENT_W, Inches(2.2))
    tf = bio.text_frame
    tf.word_wrap = True
    p = clear_tf(tf)
    run(p, "The rail is the productisation of work he has already run for CRM orgs — not a thesis from outside the stack. Small team across Ahmedabad, Bengaluru, and Jaipur. This round hires engineering, customer success, and partner sales to scale that motion.",
        size=16, color=MUTED_INK, font=FONT_BODY)
    add_footer(s, dark=False)

    # 12 ASK
    s = new_slide(prs, True)
    add_logo(s, dark=True)
    tag = tb(s, Inches(10.2), Inches(0.34), Inches(2.55), Inches(0.3))
    p = clear_tf(tag.text_frame)
    p.alignment = PP_ALIGN.RIGHT
    run(p, "THE ASK", size=11, color=AMBER, font=FONT_MONO)
    add_title(s, "₹3–5 Cr to put the rail in front of more CRM orgs.", size=28, y=Inches(0.95), h=Inches(0.7))
    add_lead(s, "The products exist. The buyer is named. This round is distribution, certifications, and the team to run more of the same — not a rebuild.",
             y=Inches(1.7), h=Inches(0.7))
    add_round(s, ML, Inches(2.55), CONTENT_W, Inches(2.55), INK_SOFT, AMBER, adj=0.04)
    hx = tb(s, ML + Inches(0.32), Inches(2.68), Inches(11.4), Inches(0.45))
    p = clear_tf(hx.text_frame)
    run(p, "Use of funds", size=24, color=PAPER, font=FONT_DISPLAY)
    funds = tb(s, ML + Inches(0.32), Inches(3.2), Inches(5.6), Inches(1.7))
    tf = funds.text_frame
    tf.word_wrap = True
    p = para(tf, "", first=True, space_after=10)
    run(p, "GTM  ", size=15, bold=True, color=PAPER, font=FONT_BODY)
    run(p, "Dreamforce, SI partners, customer success on the named buyer", size=15, color=MUTED_LIGHT, font=FONT_BODY)
    p = para(tf, "", space_after=0)
    run(p, "Product  ", size=15, bold=True, color=PAPER, font=FONT_BODY)
    run(p, "EMIE, connectors, mapping depth on the rail that already ships", size=15, color=MUTED_LIGHT, font=FONT_BODY)
    funds2 = tb(s, Inches(7.15), Inches(3.2), Inches(5.4), Inches(1.7))
    tf = funds2.text_frame
    tf.word_wrap = True
    p = para(tf, "", first=True, space_after=10)
    run(p, "Trust  ", size=15, bold=True, color=PAPER, font=FONT_BODY)
    run(p, "SOC 2, ISO 27001, GDPR / DPDP", size=15, color=MUTED_LIGHT, font=FONT_BODY)
    p = para(tf, "", space_after=0)
    run(p, "Team  ", size=15, bold=True, color=PAPER, font=FONT_BODY)
    run(p, "engineering, partnerships, delivery", size=15, color=MUTED_LIGHT, font=FONT_BODY)

    contacts = [
        ("inquire@pravadhlabs.com", "mailto:inquire@pravadhlabs.com"),
        ("+91 99985 25192", "tel:+919998525192"),
        ("Calendly", "https://calendly.com/partners-pravadhlabs/30min"),
        ("pravadhlabs.com", "https://pravadhlabs.com"),
    ]
    x = ML
    for label, url in contacts:
        box = tb(s, x, Inches(5.35), Inches(3.0), Inches(0.4))
        tf = box.text_frame
        p = clear_tf(tf)
        r = run(p, label, size=15, color=AMBER, font=FONT_BODY)
        add_hyperlink(r, url)
        x += Inches(3.05)
    legal = tb(s, ML, Inches(5.85), CONTENT_W, Inches(0.7))
    tf = legal.text_frame
    tf.word_wrap = True
    p = clear_tf(tf)
    run(p, "Pravadh Labs Private Limited · Ahmedabad · Bengaluru · Jaipur · Instrument and dilution in conversation.",
        size=14, color=MUTED_LIGHT, font=FONT_BODY)
    add_footer(s, "Confidential", "12 / 12", dark=True)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT))
    if DESKTOP.exists():
        prs.save(str(DESKTOP_OUT))
    return OUT


if __name__ == "__main__":
    path = build()
    print(f"Created: {path} ({path.stat().st_size} bytes)")
    if DESKTOP_OUT.exists():
        print(f"Copied:  {DESKTOP_OUT}")
